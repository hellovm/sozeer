"""导出：把一本书渲染成 TXT / Markdown / EPUB。

导出是"成品交付"的最后一步——前端的"原始数据呈现、拒绝美化"原则在这里同样适用：
导出的就是库里的正文，不做任何二创。EPUB 用标准库 zipfile 生成，无需第三方依赖。
"""
from __future__ import annotations

import io
import re
import zipfile
from html import escape
from pathlib import Path
from typing import Any

from apps.api.models import Book, Chapter


def _chapters(db, book_id: str) -> list[Chapter]:
    from sqlalchemy import select
    return db.scalars(
        select(Chapter).where(Chapter.book_id == book_id).order_by(Chapter.idx)
    ).all()


def to_txt(book: Book, chapters: list[Chapter]) -> str:
    parts = []
    if book.title:
        parts.append(book.title)
    if book.author:
        parts.append(f"作者：{book.author}")
    parts.append("")
    for c in chapters:
        parts.append(f"第{c.idx:03d}章 {c.title}")
        parts.append("")
        parts.append(c.content or "")
        parts.append("")
        parts.append("—" * 20)
        parts.append("")
    return "\n".join(parts)


def to_md(book: Book, chapters: list[Chapter]) -> str:
    parts = []
    parts.append(f"# {book.title}")
    if book.author:
        parts.append(f"> 作者：{book.author}")
    if book.synopsis:
        parts.append("")
        parts.append(f"> {book.synopsis}")
    parts.append("")
    for c in chapters:
        parts.append(f"## 第{c.idx:03d}章 {c.title}")
        parts.append("")
        parts.append(c.content or "")
        parts.append("")
    return "\n".join(parts)


def _sanitize(fn: str) -> str:
    return re.sub(r"[^\w一-龥\- ]", "_", fn).strip() or "book"


def to_epub(book: Book, chapters: list[Chapter]) -> bytes:
    buf = io.BytesIO()
    title = book.title or "小说"
    author = book.author or "未知"
    fname = _sanitize(title)[:60]

    spine_items = []
    xhtml_parts = []
    for c in chapters:
        cid = f"ch{c.idx:03d}"
        paras = "".join(f"<p>{escape(p)}</p>" for p in (c.content or "").splitlines() if p.strip())
        xhtml_parts.append(
            f'<?xml version="1.0" encoding="utf-8"?>\n'
            f'<html xmlns="http://www.w3.org/1999/xhtml">\n<head>\n'
            f'<meta charset="utf-8"/><title>{escape(c.title or cid)}</title>\n</head>\n'
            f'<body>\n<h1>第{c.idx:03d}章 {escape(c.title or "")}</h1>\n{paras}\n</body>\n</html>'
        )
        spine_items.append((cid, f"OEBPS/{cid}.xhtml"))

    # 目录
    nav = (
        '<?xml version="1.0" encoding="utf-8"?>\n'
        '<html xmlns="http://www.w3.org/1999/xhtml" '
        'xmlns:epub="http://www.idpf.org/2007/ops">\n<head><meta charset="utf-8"/>'
        f'<title>{escape(title)}</title></head>\n<body>\n<nav epub:type="toc" id="toc">\n'
        f'<h1>目录</h1>\n<ol>\n'
        + "".join(f'<li><a href="{fn}">{escape("第%s章 %s" % (c.idx, c.title))}</a></li>'
                  for c, (_, fn) in zip(chapters, spine_items))
        + "</ol>\n</nav>\n</body>\n</html>"
    )

    manifest = "".join(
        f'<item id="{cid}" href="{fn}" media-type="application/xhtml+xml"/>'
        for cid, fn in spine_items
    )
    manifest += '<item id="nav" href="OEBPS/nav.xhtml" media-type="application/xhtml+xml" properties="nav"/>'
    spine = "".join(f'<itemref idref="{cid}"/>' for cid, _ in spine_items)

    opf = (
        '<?xml version="1.0" encoding="utf-8"?>\n'
        '<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="bookid">\n'
        f'<metadata xmlns:dc="http://purl.org/dc/elements/1.1/">\n'
        '<dc:identifier id="bookid">urn:uuid:novelforge-' + fname + '</dc:identifier>\n'
        f'<dc:title>{escape(title)}</dc:title>\n'
        f'<dc:creator>{escape(author)}</dc:creator>\n'
        '<dc:language>zh-CN</dc:language>\n</metadata>\n'
        "<manifest>\n" + manifest + "</manifest>\n"
        "<spine>\n" + spine + "</spine>\n</package>"
    )

    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        # mimetype 必须是第一个且未压缩
        zi = zipfile.ZipInfo("mimetype")
        zi.compress_type = zipfile.ZIP_STORED
        z.writestr(zi, "application/epub+zip")
        z.writestr("META-INF/container.xml",
                   '<?xml version="1.0"?>\n'
                   '<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">\n'
                   '<rootfiles><rootfile full-path="OEBPS/content.opf" '
                   'media-type="application/oebps-package+xml"/></rootfiles>\n</container>')
        z.writestr("OEBPS/content.opf", opf)
        z.writestr("OEBPS/nav.xhtml", nav)
        for (cid, fn), xhtml in zip(spine_items, xhtml_parts):
            z.writestr(fn, xhtml)
    return buf.getvalue()


def _docx_lazy():
    """延迟导入 python-docx（★ 铁律：能模块级导入就别延迟——这里是**可选依赖**
    的例外：老部署环境没装 docx 时，txt/md/epub 导出必须照常工作）。
    未安装时抛带安装指引的 ImportError。"""
    try:
        import docx                                    # noqa: F401
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.shared import Pt
        return docx, WD_ALIGN_PARAGRAPH, qn, Pt
    except ImportError as e:
        raise ImportError(
            "DOCX 导出需要 python-docx：pip install python-docx") from e


def to_docx(book: Book, chapters: list[Chapter]) -> bytes:
    """DOCX 导出（v1.0-E4.11）：网文平台投稿普遍要 docx。

    结构：书名（H1）→ 作者/简介 → 每章 H2 + 正文段落。
    ★ 中文字体：样式里显式设 eastAsia 字体，否则 Word 显示宋体回退混乱。
    """
    docx, WD_ALIGN_PARAGRAPH, qn, Pt = _docx_lazy()
    from docx import Document

    d = Document()
    # 正文样式设中文字体（eastAsia），西文给 Calibri
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(12)
    st._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")  # type: ignore[union-attr]

    d.add_heading(book.title or "未命名作品", level=1)
    if book.author:
        p = d.add_paragraph(f"作者：{book.author}")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if book.synopsis:
        p = d.add_paragraph(book.synopsis)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for c in chapters:
        d.add_heading(f"第{c.idx:03d}章 {c.title or ''}", level=2)
        for para in (c.content or "").splitlines():
            if para.strip():
                d.add_paragraph(para.strip())
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def to_pdf(book: Book, chapters: list[Chapter]) -> bytes:
    """PDF 导出（v1.0-E4.11）：reportlab + CID 内置中文字体。

    ★ 为什么用 STSong-Light（CID 字体）而不是注册 TTF：
      服务器（尤其 Linux 容器）**没有中文字体文件**，注册 TTF 会直接
      FileNotFoundError。CID 字体内置在 reportlab 的 CMap 包里，
      零系统依赖 —— 部署安全（部署铁律：不依赖宿主机字体）。
    """
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import mm
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    except ImportError as e:
        raise ImportError(
            "PDF 导出需要 reportlab：pip install reportlab") from e

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    title_st = ParagraphStyle("T", fontName="STSong-Light", fontSize=22,
                              leading=30, alignment=1, spaceAfter=8)
    meta_st = ParagraphStyle("M", fontName="STSong-Light", fontSize=11,
                             leading=16, alignment=1, textColor="#555555")
    h2_st = ParagraphStyle("H2", fontName="STSong-Light", fontSize=15,
                           leading=22, spaceBefore=14, spaceAfter=6)
    body_st = ParagraphStyle("B", fontName="STSong-Light", fontSize=11,
                             leading=19, firstLineIndent=22)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4,
                            topMargin=20 * mm, bottomMargin=20 * mm,
                            leftMargin=22 * mm, rightMargin=22 * mm)
    story: list = [Paragraph(escape(book.title or "未命名作品"), title_st)]
    if book.author:
        story.append(Paragraph(escape(f"作者：{book.author}"), meta_st))
    if book.synopsis:
        story.append(Spacer(1, 6))
        story.append(Paragraph(escape(book.synopsis), meta_st))
    story.append(Spacer(1, 14))
    for c in chapters:
        story.append(Paragraph(escape(f"第{c.idx:03d}章 {c.title or ''}"), h2_st))
        for para in (c.content or "").splitlines():
            if para.strip():
                story.append(Paragraph(escape(para.strip()), body_st))
    doc.build(story)
    return buf.getvalue()


def to_submission(book: Book, chapters: list[Chapter], meta: dict | None = None, bible: dict | None = None) -> str:
    """投稿包：立项信息 + 第N章正文 + 创作设定集（投稿备查）。

    仅组织“可直接用于投稿准备”的成品结构，不绑定任何具体平台。
    """
    meta = meta or {}
    lines: list[str] = []
    lines.append(book.title or "未命名作品")
    if book.genre:
        lines.append(f"类型：{book.genre}")
    if meta.get("category"):
        lines.append(f"平台分类：{meta['category']}")
    if meta.get("tags"):
        lines.append("标签：" + "、".join(meta["tags"]))
    if meta.get("logline"):
        lines.append(f"一句话梗概：{meta['logline']}")
    if meta.get("selling_points"):
        lines.append("核心卖点：")
        for i, sp in enumerate(meta["selling_points"], 1):
            lines.append(f"  {i}. {sp}")
    if meta.get("synopsis"):
        lines.append("简介：")
        lines.append(meta["synopsis"])
    lines.append("")
    lines.append("=" * 24 + " 正文 " + "=" * 24)
    lines.append("")
    for c in chapters:
        lines.append(f"第{c.idx:03d}章 {c.title}")
        lines.append("")
        lines.append(c.content or "")
        lines.append("")
    # 创作设定集（投稿备查，不计入正文）
    bible = bible or {}
    if bible.get("world") or bible.get("characters") or bible.get("volumes") or bible.get("foreshadows"):
        lines.append("")
        lines.append("=" * 20 + " 创作设定集（投稿备查） " + "=" * 20)
        lines.append("")
        if bible.get("world"):
            buck: dict[str, list] = {}
            for w in bible["world"]:
                buck.setdefault(w.get("category") or "其他", []).append(w)
            for cat, items in buck.items():
                lines.append(f"【{cat}】")
                for w in items:
                    lines.append(f"  · {w.get('key')}：{w.get('value')}")
        if bible.get("characters"):
            lines.append("【人物】")
            for c in bible["characters"]:
                extra = f"（{c.get('role')}）" if c.get("role") else ""
                lines.append(f"  · {c.get('name')}{extra}：{c.get('profile') or ''}")
        if bible.get("volumes"):
            lines.append("【大纲/分卷】")
            for v in bible["volumes"]:
                lines.append(f"  · 卷{v.get('idx')}《{v.get('name')}》：{v.get('outline') or ''}")
        if bible.get("foreshadows"):
            lines.append("【伏笔】")
            for f in bible["foreshadows"]:
                lines.append(f"  · {f.get('content')}（{f.get('status')}）")
        if bible.get("signing") and not bible["signing"].get("error"):
            sg = bible["signing"]
            mark = {"pass": "✓", "warn": "△", "error": "✗"}
            lines.append("")
            lines.append("=" * 18 + f" 签约就绪度（{sg.get('platform', '')}） " + "=" * 18)
            lines.append(f"总分 {sg.get('score')}｜{sg.get('summary', '')}")
            lines.append(f"封面规范：{sg.get('cover_spec', '')}｜黄金三章：{sg.get('golden_three', '')}")
            for c in sg.get("checklist", []):
                tail = f"（{c.get('count')} 处）" if c.get("count") else ""
                ex = f"  例：{c.get('sample', '')[:40]}" if c.get("sample") else ""
                lines.append(f"  {mark.get(c.get('status'), '·')} {c.get('item')}：{c.get('status')}{tail}{ex}")
    return "\n".join(lines)


def build(book: Book, chapters: list[Chapter], fmt: str, meta: dict | None = None, bible: dict | None = None) -> tuple[bytes, str, str]:
    fmt = (fmt or "txt").lower()
    if fmt == "md":
        return to_md(book, chapters).encode("utf-8"), "text/markdown", f"{_sanitize(book.title)}.md"
    if fmt == "epub":
        return to_epub(book, chapters), "application/epub+zip", f"{_sanitize(book.title)}.epub"
    if fmt in ("publish", "submission"):
        return to_submission(book, chapters, meta, bible).encode("utf-8"), "text/plain", f"{_sanitize(book.title)}_投稿包.txt"
    if fmt == "txt":
        return to_txt(book, chapters).encode("utf-8"), "text/plain", f"{_sanitize(book.title)}.txt"
    if fmt == "docx":
        return to_docx(book, chapters), ("application/vnd.openxmlformats-officedocument"
                                         ".wordprocessingml.document"), f"{_sanitize(book.title)}.docx"
    if fmt == "pdf":
        return to_pdf(book, chapters), "application/pdf", f"{_sanitize(book.title)}.pdf"
    # 未知格式必须显式报错。原来是 fallthrough 到 to_txt：调用方拼错一个格式名
    # 会拿到一份纯 txt，还以为是投稿包，没有任何提示 —— 静默降级比报错危险得多。
    raise ValueError(f"不支持的导出格式：{fmt}（可用：txt / md / epub / docx / pdf / publish）")
